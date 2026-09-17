from __future__ import annotations

import copy
import shutil
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from product_variant_resolver.release_staging import (
    HEADERS,
    INPUTS,
    build_snapshot,
    check_bundle,
    parse_workbook,
    publish_bundle,
    validate_snapshot,
)
from tests.release_staging_test_support import write_exact_synthetic_inputs, write_workbook

ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture(scope="module")
def synthetic_root(tmp_path_factory: pytest.TempPathFactory) -> Path:
    root = tmp_path_factory.mktemp("release-staging-synthetic")
    write_exact_synthetic_inputs(root)
    (root / "data" / "external" / "hot-wheels-wiki").mkdir(parents=True)
    return root


@pytest.fixture(scope="module")
def snapshot(synthetic_root: Path) -> dict[str, Any]:
    return build_snapshot(synthetic_root)


@pytest.fixture(scope="module")
def local_snapshot() -> dict[str, Any]:
    missing = [
        filename
        for filename, _count in INPUTS.values()
        if not (ROOT / "HW data" / filename).is_file()
    ]
    if missing:
        pytest.skip(
            "owner-supplied HW data is not present; skipping exact 1,763-row local integration"
        )
    return build_snapshot(ROOT)


def test_four_local_workbooks_build_exact_review_only_snapshot(
    local_snapshot: dict[str, Any],
) -> None:
    counts = local_snapshot["counts"]
    assert isinstance(counts, dict)
    assert counts == {
        "staged_observations": 1763,
        "unique_source_record_ids": 1763,
        "unique_toy_numbers": 1763,
        "unique_casting_names": 678,
        "unknown_colors": 1763,
        "variant_note_observations": 699,
        "parse_errors": 0,
        "reviewed_variants": 0,
        "canonical_products": 0,
        "postgresql_rows_written": 0,
        "network_requests": 0,
    }
    records = local_snapshot["records"]
    assert isinstance(records, list) and len(records) == 1763
    assert all(record["color"] is None for record in records)
    assert all(record["canonical_uuid"] is None for record in records)
    assert all(record["review_status"] == "needs_canonical_review" for record in records)
    assert all(record["usage"] == "staging_only_not_evaluation_or_canonical" for record in records)


def test_local_provenance_and_raw_fields_are_preserved(local_snapshot: dict[str, Any]) -> None:
    records = local_snapshot["records"]
    assert isinstance(records, list)
    first = records[0]
    assert first["input_filename"] == "catalog-2023.xlsx"
    assert len(first["input_sha256"]) == 64
    assert first["source_record_id"] == "selenium-row-9e7c66de1126fef2"
    assert first["toy_number"] == "HKG27"
    assert first["source_row"] == 1
    assert first["raw_fields"]["Toy #"] == "HKG27"


def test_same_bytes_produce_same_snapshot(
    snapshot: dict[str, Any],
    synthetic_root: Path,
) -> None:
    second = build_snapshot(synthetic_root)
    assert second == snapshot
    assert second["batch_id"] == "local-release-staging-v1-" + second["content_sha256"]


def test_bundle_roundtrip_is_exact(snapshot: dict[str, Any], synthetic_root: Path) -> None:
    output = synthetic_root / "data/external/hot-wheels-wiki/local-export-2023-2026"
    assert publish_bundle(output, synthetic_root) == snapshot
    assert check_bundle(output, synthetic_root) == snapshot


def test_changed_normalized_record_is_rejected(
    snapshot: dict[str, Any],
    synthetic_root: Path,
) -> None:
    changed = copy.deepcopy(snapshot)
    changed["records"][0]["color"] = "red"
    with pytest.raises(ValueError, match="complete local inputs"):
        validate_snapshot(changed, synthetic_root)


def test_formula_in_release_data_is_rejected(tmp_path: Path) -> None:
    path = write_workbook(tmp_path, year=2023, count=2)
    workbook = load_workbook(path)
    workbook["Releases"]["D7"] = "=1+1"
    workbook.save(path)
    with pytest.raises(ValueError, match="formula cell"):
        parse_workbook(path, year=2023, expected_count=2, root=tmp_path)


def test_header_change_is_rejected(tmp_path: Path) -> None:
    path = write_workbook(tmp_path, year=2024, count=2)
    workbook = load_workbook(path)
    workbook["Releases"]["A6"] = "Different ID"
    workbook.save(path)
    with pytest.raises(ValueError, match="header differs"):
        parse_workbook(path, year=2024, expected_count=2, root=tmp_path)
    assert len(HEADERS) == 19


def test_extra_twentieth_column_value_is_rejected(tmp_path: Path) -> None:
    path = write_workbook(tmp_path, year=2023, count=2)
    workbook = load_workbook(path)
    workbook["Releases"]["T6"] = "Unexpected column"
    workbook["Releases"]["T7"] = "silently ignored before this regression test"
    workbook.save(path)
    with pytest.raises(ValueError, match=r"extra column content.*T6"):
        parse_workbook(path, year=2023, expected_count=2, root=tmp_path)


def test_extra_twentieth_column_formula_is_rejected(tmp_path: Path) -> None:
    path = write_workbook(tmp_path, year=2023, count=2)
    workbook = load_workbook(path)
    workbook["Releases"]["T7"] = "=1+1"
    workbook.save(path)
    with pytest.raises(ValueError, match=r"formula cell.*T7"):
        parse_workbook(path, year=2023, expected_count=2, root=tmp_path)


def test_noncontiguous_source_rows_are_rejected(tmp_path: Path) -> None:
    path = write_workbook(tmp_path, year=2023, count=2)
    workbook = load_workbook(path)
    workbook["Releases"]["O8"] = 3
    workbook.save(path)
    with pytest.raises(ValueError, match="source rows must be ordered and contiguous"):
        parse_workbook(path, year=2023, expected_count=2, root=tmp_path)


def test_duplicate_toy_number_is_rejected(tmp_path: Path, synthetic_root: Path) -> None:
    shutil.copytree(synthetic_root / "HW data", tmp_path / "HW data")
    path = tmp_path / "HW data" / "catalog-2025.xlsx"
    workbook = load_workbook(path)
    workbook["Releases"]["D8"] = workbook["Releases"]["D7"].value
    workbook.save(path)
    with pytest.raises(ValueError, match="duplicate toy numbers"):
        build_snapshot(tmp_path)


def test_duplicate_source_record_id_is_rejected(tmp_path: Path, synthetic_root: Path) -> None:
    shutil.copytree(synthetic_root / "HW data", tmp_path / "HW data")
    path = tmp_path / "HW data" / "catalog-2026.xlsx"
    workbook = load_workbook(path)
    workbook["Releases"]["A8"] = workbook["Releases"]["A7"].value
    workbook.save(path)
    with pytest.raises(ValueError, match="duplicate source record IDs"):
        build_snapshot(tmp_path)
