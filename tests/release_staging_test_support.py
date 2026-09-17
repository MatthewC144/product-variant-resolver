from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook

from product_variant_resolver.release_staging import HEADERS, INPUTS


def write_workbook(root: Path, *, year: int, count: int, global_offset: int = 0) -> Path:
    directory = root / "HW data"
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"catalog-{year}.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Releases"
    for column, header in enumerate(HEADERS, start=1):
        sheet.cell(row=6, column=column, value=header)
    for source_row in range(1, count + 1):
        global_index = global_offset + source_row - 1
        casting = f"Synthetic Casting {global_index % 678:03d}"
        toy_number = f"SYN-{year}-{source_row:04d}"
        raw_fields = json.dumps(
            {
                "Col.#": str(source_row),
                "Model Name": casting,
                "Photo": "",
                "Series": "Synthetic Series",
                "Series #": "1/10",
                "Toy #": toy_number,
            },
            sort_keys=True,
        )
        values = (
            f"synthetic-row-{year}-{source_row:04d}",
            year,
            "Hot Wheels",
            toy_number,
            str(source_row),
            casting,
            casting,
            None,
            "Synthetic Series",
            "1/10",
            None,
            f"Synthetic {year} source",
            f"https://example.invalid/{year}",
            1,
            source_row,
            "parsed",
            None,
            "2026-09-17T00:00:00Z",
            raw_fields,
        )
        for column, value in enumerate(values, start=1):
            sheet.cell(row=source_row + 6, column=column, value=value)
    workbook.save(path)
    return path


def write_exact_synthetic_inputs(root: Path) -> None:
    offset = 0
    for year, (_filename, count) in INPUTS.items():
        write_workbook(root, year=year, count=count, global_offset=offset)
        offset += count
