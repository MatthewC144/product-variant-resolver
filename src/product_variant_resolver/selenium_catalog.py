from __future__ import annotations

import hashlib
import json
import re
from datetime import UTC, datetime
from pathlib import Path
from typing import TYPE_CHECKING, Any, cast
from urllib.parse import urlparse

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

DENIED_HOST_SUFFIXES = ("fandom.com", "wikia.org")
REQUIRED_HEADER_GROUPS = (
    frozenset({"model", "modelname", "casting", "castingname", "name"}),
    frozenset(
        {"toy", "toy#", "toynumber", "toyno", "collector", "collector#", "collectornumber", "col#"}
    ),
)
HEADER_ALIASES = {
    "toy": "toy_number",
    "toy#": "toy_number",
    "toynumber": "toy_number",
    "toyno": "toy_number",
    "modelnumber": "toy_number",
    "collector": "collector_number",
    "collector#": "collector_number",
    "collectornumber": "collector_number",
    "collectorno": "collector_number",
    "col#": "collector_number",
    "model": "source_model_label",
    "modelname": "source_model_label",
    "casting": "source_model_label",
    "castingname": "source_model_label",
    "name": "source_model_label",
    "series": "series",
    "seriesname": "series",
    "series#": "series_position",
    "seriesnumber": "series_position",
    "seriesno": "series_position",
    "position": "series_position",
    "color": "color",
    "colour": "color",
    "releaseyear": "release_year",
    "year": "release_year",
    "variant": "variant_note",
    "notes": "variant_note",
    "note": "variant_note",
}
CHALLENGE_MARKERS = (
    "verify you are human",
    "captcha",
    "access denied",
    "too many requests",
    "rate limit exceeded",
    "checking your browser",
)


TABLE_EXTRACTION_SCRIPT = r"""
const clean = (value) => (value || "").replace(/\s+/g, " ").trim();
const candidates = Array.from(document.querySelectorAll(
  "table.wikitable, table.article-table, table.fandom-table, .mw-parser-output table"
));
return candidates.map((table, tableIndex) => {
  const allRows = Array.from(table.querySelectorAll("tr"));
  let headerIndex = allRows.findIndex((row) => row.querySelectorAll("th").length > 0);
  if (headerIndex < 0) headerIndex = 0;
  const headers = Array.from(allRows[headerIndex]?.querySelectorAll("th,td") || [])
    .map((cell) => clean(cell.innerText || cell.textContent));
  const rows = allRows.slice(headerIndex + 1).map((row, rowOffset) => {
    const cells = Array.from(row.querySelectorAll(":scope > th, :scope > td")).map((cell) => ({
      text: clean(cell.innerText || cell.textContent),
      links: Array.from(cell.querySelectorAll("a[href]")).map((link) => link.href),
    }));
    return {row_index: rowOffset + 1, cells};
  }).filter((row) => row.cells.some((cell) => cell.text));
  return {table_index: tableIndex + 1, headers, rows};
});
"""


def _normalized_header(value: str) -> str:
    value = value.casefold().replace("№", "number")
    value = re.sub(r"\bno\.?\b", "no", value)
    return re.sub(r"[^a-z0-9#]+", "", value)


def _normalized_host(value: str) -> str:
    return value.strip().casefold().rstrip(".")


def is_denied_host(host: str) -> bool:
    normalized = _normalized_host(host)
    return any(
        normalized == denied or normalized.endswith(f".{denied}") for denied in DENIED_HOST_SUFFIXES
    )


def validate_remote_target(start_url: str, allowed_host: str) -> str:
    parsed = urlparse(start_url)
    host = _normalized_host(parsed.hostname or "")
    allowed = _normalized_host(allowed_host)
    if parsed.scheme not in {"http", "https"} or not host:
        raise ValueError("start URL must be an absolute http(s) URL")
    if not allowed or is_denied_host(allowed) or is_denied_host(host):
        raise ValueError("Fandom and Wikia hosts are intentionally blocked")
    if host != allowed:
        raise ValueError("start URL host must exactly match --allowed-host")
    if parsed.username or parsed.password:
        raise ValueError("credentials must not be embedded in the start URL")
    return host


def _variant_parts(model_label: str) -> tuple[str, str | None]:
    match = re.search(
        r"\s*\(((?:\d+(?:st|nd|rd|th)\s+)?Color(?:\s*-\s*[^)]+)?)\)\s*$",
        model_label,
        flags=re.IGNORECASE,
    )
    if not match:
        return model_label.strip(), None
    return model_label[: match.start()].strip(), match.group(1).strip()


def _canonical_headers(headers: list[str]) -> list[str | None]:
    return [HEADER_ALIASES.get(_normalized_header(header)) for header in headers]


def _eligible_table(headers: list[str]) -> bool:
    normalized = {_normalized_header(header) for header in headers}
    return all(normalized.intersection(group) for group in REQUIRED_HEADER_GROUPS)


def normalize_table_snapshots(
    snapshots: list[dict[str, Any]],
    *,
    source_url: str,
    page_title: str,
    release_year: int,
    brand: str,
    scraped_at_utc: str,
    max_records: int,
) -> tuple[list[dict[str, Any]], dict[str, int]]:
    if not 1900 <= release_year <= 2100:
        raise ValueError("release year must be between 1900 and 2100")
    if not 1 <= max_records <= 5_000:
        raise ValueError("max records must be between 1 and 5000")

    records: list[dict[str, Any]] = []
    eligible_tables = 0
    rows_seen = 0
    rows_filtered_by_year = 0
    for table in snapshots:
        headers = [str(value) for value in table.get("headers", [])]
        if not _eligible_table(headers):
            continue
        eligible_tables += 1
        canonical_headers = _canonical_headers(headers)
        table_index = int(table.get("table_index", eligible_tables))
        for raw_row in table.get("rows", []):
            rows_seen += 1
            cells = raw_row.get("cells", [])
            raw_fields: dict[str, str] = {}
            normalized: dict[str, str] = {}
            for index, header in enumerate(headers):
                text = ""
                if index < len(cells) and isinstance(cells[index], dict):
                    text = str(cells[index].get("text", "")).strip()
                raw_fields[header or f"column_{index + 1}"] = text
                canonical = canonical_headers[index]
                if canonical and text and canonical not in normalized:
                    normalized[canonical] = text

            row_year_text = normalized.get("release_year", "")
            if row_year_text:
                match = re.search(r"(?:19|20)\d{2}", row_year_text)
                if match and int(match.group()) != release_year:
                    rows_filtered_by_year += 1
                    continue

            model_label = normalized.get("source_model_label", "").strip()
            casting_name, inferred_variant = _variant_parts(model_label)
            variant_note = normalized.get("variant_note") or inferred_variant
            errors: list[str] = []
            if not model_label:
                errors.append("missing model/casting name")
            if not normalized.get("toy_number") and not normalized.get("collector_number"):
                errors.append("missing toy and collector numbers")
            source_row = int(raw_row.get("row_index", rows_seen))
            record_key = f"{source_url}:{table_index}:{source_row}:{model_label}"
            records.append(
                {
                    "source_record_id": "selenium-row-"
                    + hashlib.sha256(record_key.encode("utf-8")).hexdigest()[:16],
                    "release_year": release_year,
                    "brand": brand,
                    "toy_number": normalized.get("toy_number") or None,
                    "collector_number": normalized.get("collector_number") or None,
                    "source_model_label": model_label or None,
                    "casting_name": casting_name or None,
                    "variant_note": variant_note,
                    "series": normalized.get("series") or None,
                    "series_position": normalized.get("series_position") or None,
                    "color": normalized.get("color") or None,
                    "source_page_title": page_title,
                    "source_page_url": source_url,
                    "source_table_index": table_index,
                    "source_row": source_row,
                    "parse_status": "error" if errors else "parsed",
                    "parse_error": "; ".join(errors) or None,
                    "scraped_at_utc": scraped_at_utc,
                    "raw_fields_json": json.dumps(raw_fields, ensure_ascii=False, sort_keys=True),
                }
            )
            if len(records) >= max_records:
                break
        if len(records) >= max_records:
            break

    if eligible_tables == 0:
        raise ValueError("no catalog table with model and identifier headers was found")
    if not records:
        raise ValueError(f"no rows matched release year {release_year}")
    return records, {
        "candidate_tables": len(snapshots),
        "eligible_tables": eligible_tables,
        "rows_seen": rows_seen,
        "rows_filtered_by_year": rows_filtered_by_year,
        "records_exported": len(records),
    }


def create_chrome_driver(*, headless: bool, page_load_timeout: float) -> Any:
    try:
        from selenium import webdriver
        from selenium.webdriver.chrome.options import Options
    except ImportError as error:
        raise RuntimeError(
            "Selenium is not installed. Run: uv pip install --python .venv/bin/python -e '.[crawler]'"
        ) from error

    options = Options()
    if headless:
        options.add_argument("--headless=new")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--window-size=1440,1200")
    options.add_experimental_option(
        "prefs",
        {
            "profile.managed_default_content_settings.images": 2,
            "download_restrictions": 3,
        },
    )
    driver = webdriver.Chrome(options=options)
    driver.set_page_load_timeout(page_load_timeout)
    return driver


def collect_page(
    driver: Any,
    *,
    start_url: str,
    allowed_host: str,
    release_year: int,
    brand: str,
    wait_timeout: float,
    max_records: int,
) -> dict[str, Any]:
    validate_remote_target(start_url, allowed_host)
    driver.get(start_url)
    final_url = str(driver.current_url)
    final_host = _normalized_host(urlparse(final_url).hostname or "")
    if final_host != _normalized_host(allowed_host) or is_denied_host(final_host):
        raise RuntimeError("navigation left the exact allowed host; collection stopped")

    try:
        from selenium.webdriver.common.by import By
        from selenium.webdriver.support.ui import WebDriverWait
    except ImportError as error:
        raise RuntimeError("Selenium is required to collect a page") from error

    WebDriverWait(driver, wait_timeout).until(
        lambda browser: browser.find_elements(
            By.CSS_SELECTOR,
            "table.wikitable, table.article-table, table.fandom-table, .mw-parser-output table",
        )
    )
    body_text = str(driver.find_element(By.TAG_NAME, "body").text).casefold()
    title = str(driver.title).strip()
    challenge_text = f"{title}\n{body_text[:4000]}".casefold()
    if any(marker in challenge_text for marker in CHALLENGE_MARKERS):
        raise RuntimeError("access challenge or rate-limit page detected; collection stopped")

    snapshots = driver.execute_script(TABLE_EXTRACTION_SCRIPT)
    if not isinstance(snapshots, list):
        raise TypeError("browser returned an unexpected table snapshot")
    scraped_at = datetime.now(UTC).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    records, counters = normalize_table_snapshots(
        snapshots,
        source_url=final_url,
        page_title=title,
        release_year=release_year,
        brand=brand,
        scraped_at_utc=scraped_at,
        max_records=max_records,
    )
    return {
        "schema_version": "pvr-selenium-catalog-export-v1",
        "source": {
            "requested_url": start_url,
            "final_url": final_url,
            "allowed_host": final_host,
            "page_title": title,
            "permission_confirmed_by_operator": True,
            "fandom_wikia_blocked": True,
        },
        "collection": {
            "release_year": release_year,
            "brand": brand,
            "scraped_at_utc": scraped_at,
            **counters,
        },
        "records": records,
    }


def export_xlsx(
    payload: dict[str, Any],
    *,
    output_path: Path,
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.workbook.properties import CalcProperties
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as error:
        raise RuntimeError(
            "OpenPyXL is not installed. Run: uv pip install --python .venv/bin/python -e '.[crawler]'"
        ) from error

    records = payload.get("records")
    if not isinstance(records, list) or not records:
        raise ValueError("payload must contain at least one record")
    source = payload.get("source")
    collection = payload.get("collection")
    if not isinstance(source, dict) or not isinstance(collection, dict):
        raise TypeError("payload source and collection must be objects")

    columns = (
        ("Source record ID", "source_record_id"),
        ("Release year", "release_year"),
        ("Brand", "brand"),
        ("Toy number", "toy_number"),
        ("Collector number", "collector_number"),
        ("Source model label", "source_model_label"),
        ("Casting name", "casting_name"),
        ("Variant note", "variant_note"),
        ("Series", "series"),
        ("Series position", "series_position"),
        ("Color", "color"),
        ("Source page title", "source_page_title"),
        ("Source page URL", "source_page_url"),
        ("Source table", "source_table_index"),
        ("Source row", "source_row"),
        ("Parse status", "parse_status"),
        ("Parse error", "parse_error"),
        ("Collected at (UTC)", "scraped_at_utc"),
        ("Raw fields (JSON)", "raw_fields_json"),
    )
    dark_blue = "1F4E78"
    light_blue = "DCE6F1"
    white = "FFFFFF"
    body_color = "1F2937"
    thin_gray = Side(style="thin", color="D8DEE6")

    workbook = Workbook()
    workbook.calculation = CalcProperties(
        calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True
    )
    summary = cast("Worksheet", workbook.active)
    summary.title = "Summary"
    releases = cast("Worksheet", workbook.create_sheet("Releases"))
    for sheet in (summary, releases):
        sheet.sheet_view.showGridLines = False

    summary["A2"] = "Catalog collection summary"
    summary["A2"].font = Font(name="Arial", size=14, bold=True, color=body_color)
    summary["A4"] = "Metric"
    summary["B4"] = "Value"
    metric_rows = (
        ("Source page", source.get("final_url")),
        ("Allowed host", source.get("allowed_host")),
        ("Release year", collection.get("release_year")),
        ("Exported records", f"=COUNTA(Releases!A7:A{len(records) + 6})"),
        ("Rows with parse errors", f'=COUNTIF(Releases!P7:P{len(records) + 6},"error")'),
        ("Candidate tables", collection.get("candidate_tables")),
        ("Eligible catalog tables", collection.get("eligible_tables")),
        ("Rows inspected", collection.get("rows_seen")),
        ("Rows filtered by year", collection.get("rows_filtered_by_year")),
    )
    for row_number, (label, value) in enumerate(metric_rows, start=5):
        summary.cell(row_number, 1, label)
        summary.cell(row_number, 2, value)
    for cell in summary[4]:
        cell.fill = PatternFill("solid", fgColor=dark_blue)
        cell.font = Font(name="Arial", size=10, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in summary.iter_rows(min_row=5, max_row=13, min_col=1, max_col=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color=body_color)
            cell.border = Border(bottom=thin_gray)
            cell.alignment = Alignment(vertical="center")
        row[0].fill = PatternFill("solid", fgColor=light_blue)
    summary["A15"] = "Use"
    summary["B15"] = "Review-only collection output. It is not canonical catalog data."
    summary["A16"] = "Safety"
    summary["B16"] = "The collector handles one allowlisted host and refuses Fandom/Wikia."
    for coordinate in ("A15", "A16"):
        summary[coordinate].font = Font(name="Arial", size=10, bold=True, color=body_color)
    for coordinate in ("B15", "B16"):
        summary[coordinate].font = Font(name="Arial", size=10, italic=True, color="4B5563")
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 88
    summary.sheet_properties.tabColor = dark_blue

    releases["A2"] = "Collected release rows"
    releases["A2"].font = Font(name="Arial", size=14, bold=True, color=body_color)
    releases["A3"] = "Source URL"
    releases["B3"] = source.get("final_url")
    releases["A4"] = "Collected at (UTC)"
    releases["B4"] = collection.get("scraped_at_utc")
    for coordinate in ("A3", "A4"):
        releases[coordinate].font = Font(name="Arial", size=10, bold=True, color=body_color)
    for coordinate in ("B3", "B4"):
        releases[coordinate].font = Font(name="Arial", size=10, color="4B5563")
    for column_number, (label, _key) in enumerate(columns, start=1):
        cell = releases.cell(6, column_number, label)
        cell.fill = PatternFill("solid", fgColor=dark_blue)
        cell.font = Font(name="Arial", size=10, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_number, record in enumerate(records, start=7):
        if not isinstance(record, dict):
            raise TypeError("every record must be an object")
        for column_number, (_label, key) in enumerate(columns, start=1):
            cell = releases.cell(row_number, column_number, record.get(key))
            cell.font = Font(name="Arial", size=10, color=body_color)
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(bottom=Side(style="thin", color="E5E7EB"))
        for column_number in (1, 4, 5):
            releases.cell(row_number, column_number).number_format = "@"
        releases.cell(row_number, 2).number_format = "0"
        releases.cell(row_number, 14).number_format = "0"
        releases.cell(row_number, 15).number_format = "0"
        releases.cell(row_number, 18).number_format = "@"

    last_row = len(records) + 6
    table = Table(displayName="CollectedReleases", ref=f"A6:S{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    releases.add_table(table)
    releases.freeze_panes = "D7"
    releases.conditional_formatting.add(
        f"P7:P{last_row}",
        FormulaRule(  # type: ignore[no-untyped-call]
            formula=['P7="error"'],
            fill=PatternFill("solid", fgColor="FDE2E2"),
            font=Font(bold=True, color="B91C1C"),
        ),
    )
    widths = (27, 13, 15, 14, 16, 30, 28, 22, 24, 16, 16, 28, 56, 13, 12, 14, 32, 23, 56)
    for column_number, width in enumerate(widths, start=1):
        releases.column_dimensions[get_column_letter(column_number)].width = width
    releases.sheet_properties.tabColor = "5B9BD5"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
